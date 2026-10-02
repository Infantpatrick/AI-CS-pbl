import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    doc_path = r"D:\AICS PROJECT\Project_Report_PhishGuard_AI.docx"
    if not os.path.exists(doc_path):
        print(f"File not found: {doc_path}")
        return
        
    doc = Document(doc_path)
    
    # Dictionary mapping heading texts to the list of paragraphs to inject.
    expansion_content = {
        "1.1 Background": [
            "The exponential growth of email communications in the digital era has concurrently led to a massive increase in cyber threats, with phishing being the most prominent and damaging. Phishing attacks involve the deceptive use of electronic communications to trick individuals into revealing sensitive information, such as login credentials, financial details, or personal data. These attacks rely heavily on social engineering, exploiting human psychology rather than technical vulnerabilities.",
            "Traditional rule-based systems and keyword filters, which rely on static blacklists and predefined signatures, are increasingly inadequate. Modern adversaries utilize sophisticated evasion techniques, including obfuscated URLs, homoglyph attacks, and dynamic content generation. The limitations of these deterministic systems necessitate a paradigm shift towards intelligent, adaptive solutions.",
            "Natural Language Processing (NLP) and machine learning offer a robust alternative by analyzing the semantic context and intent of the email text rather than relying on isolated keywords. Specifically, transformer-based models like Bidirectional Encoder Representations from Transformers (BERT) have revolutionized NLP by capturing deep, bidirectional contextual relationships within text, making them highly suitable for detecting the nuanced and manipulative language characteristic of phishing emails.",
            "Furthermore, an effective detection system cannot rely on text analysis alone. The inclusion of URL analysis is critical, as the ultimate goal of most phishing emails is to direct the victim to a malicious external site. Therefore, integrating NLP-based text classification with static URL heuristics creates a comprehensive, multi-layered detection workflow capable of identifying sophisticated threats with high accuracy."
        ],
        "1.2 Problem Statement": [
            "Traditional email filtering mechanisms frequently fail to detect context-based phishing attempts that do not contain overt malicious payloads or blacklisted keywords. Adversaries craft emails that mimic legitimate organizational communications, utilizing social-engineering language to invoke urgency, fear, or curiosity, which easily bypasses standard security gateways.",
            "The reliance on static blacklists is a fundamental flaw, as attackers rapidly cycle through disposable domains and IP addresses. Obfuscated URLs, URL shorteners, and seemingly benign links embedded within seemingly legitimate text further complicate the detection process. Keyword-based detection yields high false-positive rates when processing legitimate business communications that naturally contain words like 'urgent' or 'invoice'.",
            "There is a pressing need for an automated, intelligent analysis system that can interpret the semantic meaning of an email and independently evaluate the structural integrity of any embedded URLs. Moreover, end-users require interpretable risk information—not just a binary block/allow decision—to make informed security choices. The absence of an integrated, user-facing system that combines advanced NLP with transparent risk assessment forms the core problem this project addresses."
        ],
        "1.6 Overview of Proposed System": [
            "The proposed PhishGuard AI system operates through a seamless, decoupled microservices architecture designed for scalability and real-time processing. The workflow begins when a user submits an email—either via manual text input or through the Gmail-integrated inbox—to the React-based frontend. This user interface provides a responsive, intuitive environment for interaction.",
            "Upon submission, the data is transmitted via RESTful API to the FastAPI backend. The backend orchestrates the entire detection pipeline. First, the email text undergoes rigorous preprocessing, including HTML stripping, special character removal, and normalization, preparing it for the machine learning layer.",
            "The core NLP engine utilizes a fine-tuned BERT (Bidirectional Encoder Representations from Transformers) model. The preprocessed text is tokenized into input IDs and attention masks, then passed through the BERT encoder. The model outputs a continuous probability score representing the likelihood that the text contains phishing semantics. Simultaneously, a URL extraction module scans the text for hyperlinks, passing them to a static analysis engine that evaluates structural heuristics such as IP usage, excessive length, and suspicious characters.",
            "Finally, the Risk Assessment Layer synthesizes the BERT probability and the URL indicators using deterministic logic to assign a final risk level (LOW, MEDIUM, or HIGH). This comprehensive evaluation is immediately returned to the frontend for visualization and permanently logged in the SQLite database for historical reporting and dashboard metrics."
        ],
        "3.1.4 Canonical Text Preprocessing": [
            "Canonical text preprocessing is a critical prerequisite for transformer-based inference. Once the raw email is stripped of HTML and special characters, the text must be formatted into a tensor representation that the BERT model can ingest. This is achieved using the WordPiece tokenizer associated with the `bert-base-uncased` vocabulary.",
            "The tokenizer splits the normalized text into subword tokens, mapping each to a unique integer ID. To maintain structural consistency across all inputs, the sequence is either padded or truncated to a fixed maximum sequence length of 256 tokens. This length was selected to capture the core intent of the email while optimizing computational efficiency.",
            "In addition to the input IDs, the tokenizer generates an attention mask—a binary tensor indicating which tokens represent actual text and which are merely padding. This mask prevents the BERT encoder from allocating computational attention to empty padding tokens. Finally, special tokens `[CLS]` (classification) and `[SEP]` (separator) are automatically prepended and appended to the sequence, respectively, finalizing the input tensors for batch processing."
        ],
        "3.1.5 BERT Classification": [
            "The machine learning layer is powered by a pre-trained BERT (Bidirectional Encoder Representations from Transformers) model, specifically fine-tuned for binary sequence classification. Unlike traditional recurrent neural networks, BERT processes the entire sequence simultaneously using the self-attention mechanism, allowing it to evaluate the contextual significance of each word in relation to all other words in the sequence, both forwards and backwards.",
            "The architecture consists of a multi-layer bidirectional Transformer encoder. When the tokenized input IDs and attention masks are fed into the network, they are transformed through token embeddings, segment embeddings, and positional embeddings. These embeddings propagate through the encoder layers, resulting in a rich, high-dimensional contextual representation of the email.",
            "For the classification task, the model extracts the final hidden state corresponding to the `[CLS]` token. This state acts as an aggregate representation of the entire sequence. It is passed through a linear classification head consisting of a fully connected layer and a dropout layer for regularization. The output logits are then normalized using a softmax function to produce a probability distribution across the two classes: Legitimate (Class 0) and Phishing (Class 1). A classification threshold of 0.50 is utilized to determine the preliminary BERT prediction."
        ],
        "3.1.6 URL Extraction and Static Analysis": [
            "The URL Extraction and Static Analysis module operates in parallel with the NLP pipeline to evaluate the structural integrity of embedded hyperlinks. A robust regular expression engine parses the raw email body to extract all valid URLs. These URLs are then subjected to a series of static heuristic checks designed to identify common obfuscation techniques used by attackers.",
            "The static analysis evaluates multiple indicators. It checks for the presence of raw IP addresses in the hostname, which is a strong indicator of malicious intent used to bypass domain reputation filters. It analyzes the protocol, flagging HTTP usage over the secure HTTPS protocol. The length of the URL is assessed, as attackers frequently use excessively long URLs to hide malicious parameters.",
            "Furthermore, the engine counts the number of subdomains, flagging URLs that attempt to mimic legitimate domains via deep subdomain nesting (e.g., `login.paypal.com.secure.update.xyz`). It also detects the presence of the `@` symbol, which is used to ignore preceding authentication strings, and checks for excessive percent-encoding. The aggregation of these structural anomalies provides critical forensic evidence that informs the final risk assessment."
        ],
        "3.4 System Architecture": [
            "The system architecture of PhishGuard AI is designed as a decoupled, multi-tier microservices application, ensuring scalability, maintainability, and clear separation of concerns. The architecture is logically divided into several interconnected layers, each responsible for a specific domain of the detection workflow.",
            "The Presentation Layer, built with React and Tailwind CSS, provides a responsive, single-page application (SPA) interface. It handles all user interactions, input validation, and visual rendering of analysis results. The Authentication Layer secures the system using JWT (JSON Web Tokens), managing user sessions and authorization.",
            "The Backend/API Layer, implemented in FastAPI, serves as the central orchestrator. It exposes RESTful endpoints and manages data flow between the frontend and the analytical engines. The NLP/ML Layer houses the PyTorch-based BERT model, executing the heavy computational tasks required for text inference. Simultaneously, the URL Analysis Layer performs deterministic static checks on extracted hyperlinks.",
            "The Risk Assessment Layer acts as the final decision matrix, synthesizing the probabilistic output from the ML Layer and the heuristic flags from the URL Layer. Finally, the Database/History Layer utilizes SQLite to persistently store user profiles, analysis logs, and system metrics, which are later aggregated and visualized by the Reporting Layer on the frontend dashboard."
        ],
        "4.2 Modules": [
            "The implementation of PhishGuard AI is modularized to ensure high cohesion and low coupling. The Authentication Module handles secure user registration and login, utilizing bcrypt for password hashing and issuing JWTs for stateless session management.",
            "The Gmail Integration Module leverages the Google OAuth2.0 API, allowing users to securely connect their Gmail accounts. It retrieves recent emails and formats them for the Inbox Module, which provides a familiar, mail-client-like interface for users to browse and select emails for scanning.",
            "The core analytical functionality resides in the Email Analysis Module. It accepts raw email text, sanitizes it through the Data Preprocessing Module, and delegates the computational workload to the BERT Inference Module and URL Analysis Module. The Risk Assessment Module then synthesizes these discrete outputs into a unified risk profile.",
            "Post-analysis, the Reports and History Modules capture the transaction. The History Module writes the metadata, predictions, and extracted indicators to the relational database. The Dashboard Module queries this historical data to render real-time statistical visualizations, providing users with actionable intelligence regarding their threat landscape."
        ],
        "5.12 Test Cases and Results": [
            "The testing phase verified the operational integrity of all system components through rigorous, structured test cases. Unit testing ensured that isolated functions, such as URL extraction regex patterns and text normalization routines, executed flawlessly under various edge cases.",
            "API Testing was conducted using the integrated Swagger UI and Postman, verifying that the FastAPI endpoints correctly processed valid payloads, rejected malformed requests with appropriate 422 Unprocessable Entity codes, and maintained stable response times under load.",
            "End-to-End Testing simulated complete user journeys, from login to email submission, BERT inference, and historical report generation. The system consistently demonstrated accurate routing, precise threshold application, and seamless data persistence, culminating in a highly reliable application ready for deployment."
        ],
        "6.1.3 Prediction Performance Analysis": [
            "The evaluation of the BERT model provides critical insight into its predictive capabilities. The documented historical metrics establish a strong baseline for the model's performance on the validation split. An Accuracy of 67.58% indicates the overall proportion of correct classifications, while the Precision of 61.28% reflects the model's ability to minimize false positives when flagging an email as phishing.",
            "Most notably, the model achieved an exceptional Recall of 97.48%. In the context of cybersecurity and threat detection, Recall is arguably the most critical metric, as it measures the system's ability to identify true phishing emails and minimize false negatives. A high Recall ensures that the vast majority of malicious emails are successfully intercepted, prioritizing user safety.",
            "The F1-score of 75.25% represents the harmonic mean of Precision and Recall, demonstrating a robust balance in the model's predictive power. These metrics validate the efficacy of the transformer-based approach in discerning the subtle semantic indicators of social engineering, outperforming traditional deterministic keyword filters."
        ],
        "7.5 Future Enhancements": [
            "While PhishGuard AI successfully demonstrates the viability of transformer-based phishing detection, several avenues exist for future enhancement. Improved model training utilizing larger, more diverse, and continually updated datasets would further increase precision and adaptability against zero-day social engineering tactics.",
            "The URL intelligence module could be significantly upgraded by integrating dynamic analysis and real-time threat intelligence feeds (e.g., VirusTotal or Google Safe Browsing API), moving beyond static heuristics to verify domain reputation and active malicious payloads.",
            "Additionally, advancements in model explainability, such as integrating SHAP or LIME, would provide users with precise highlighting of the specific words or phrases that triggered the phishing classification. Future deployment strategies involve migrating the architecture to a containerized cloud environment using Docker and Kubernetes, ensuring high availability and enterprise-grade scalability."
        ]
    }
    
    paragraphs = doc.paragraphs
    insert_count = 0
    
    # Iterate in reverse so inserting doesn't mess up the indices
    for i in range(len(paragraphs) - 1, -1, -1):
        p = paragraphs[i]
        text = p.text.strip()
        
        # Check if this paragraph is one of our target headings
        for heading_text, paras_to_insert in expansion_content.items():
            if text.startswith(heading_text):
                # Target found. We want to insert the paragraphs AFTER this heading.
                # In reverse iteration, inserting after means we find the element and addnext in reverse order of the paras_to_insert
                for new_text in reversed(paras_to_insert):
                    new_p = doc.add_paragraph(new_text)
                    new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    for run in new_p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(12)
                    
                    # Move the new paragraph element directly after the heading
                    p._element.addnext(new_p._element)
                    insert_count += 1
                
                # Remove from dict so we only inject once
                del expansion_content[heading_text]
                break

    print(f"Successfully injected {insert_count} expansive paragraphs.")
    
    doc.save(doc_path)
    print("Report expanded and saved successfully.")

if __name__ == "__main__":
    main()
