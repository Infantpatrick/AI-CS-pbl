import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    doc_path = r"D:\AICS PROJECT\Project_Report_PhishGuard_AI.docx"
    if not os.path.exists(doc_path):
        print("File not found.")
        return
        
    doc = Document(doc_path)
    
    # Massive expansion content for 50 pages
    expansion_dict = {
        "1.1 Background": [
            "The exponential growth of email communications in the digital era has concurrently led to a massive increase in cyber threats, with phishing being the most prominent and damaging. Phishing attacks involve the deceptive use of electronic communications to trick individuals into revealing sensitive information, such as login credentials, financial details, or personal data. These attacks rely heavily on social engineering, exploiting human psychology rather than technical vulnerabilities.",
            "Traditional rule-based systems and keyword filters, which rely on static blacklists and predefined signatures, are increasingly inadequate. Modern adversaries utilize sophisticated evasion techniques, including obfuscated URLs, homoglyph attacks, and dynamic content generation. The limitations of these deterministic systems necessitate a paradigm shift towards intelligent, adaptive solutions.",
            "Natural Language Processing (NLP) and machine learning offer a robust alternative by analyzing the semantic context and intent of the email text rather than relying on isolated keywords. Specifically, transformer-based models like Bidirectional Encoder Representations from Transformers (BERT) have revolutionized NLP by capturing deep, bidirectional contextual relationships within text, making them highly suitable for detecting the nuanced and manipulative language characteristic of phishing emails.",
            "Furthermore, an effective detection system cannot rely on text analysis alone. The inclusion of URL analysis is critical, as the ultimate goal of most phishing emails is to direct the victim to a malicious external site. Therefore, integrating NLP-based text classification with static URL heuristics creates a comprehensive, multi-layered detection workflow capable of identifying sophisticated threats with high accuracy."
        ],
        "1.2 Problem Statement": [
            "Traditional email filtering mechanisms frequently fail to detect context-based phishing attempts that do not contain overt malicious payloads or blacklisted keywords. Adversaries craft emails that mimic legitimate organizational communications, utilizing social-engineering language to invoke urgency, fear, or curiosity, which easily bypasses standard security gateways.",
            "The reliance on static blacklists is a fundamental flaw, as attackers rapidly cycle through disposable domains and IP addresses. Obfuscated URLs, URL shorteners, and seemingly benign links embedded within seemingly legitimate text further complicate the detection process. Keyword-based detection yields high false-positive rates when processing legitimate business communications that naturally contain words like 'urgent' or 'invoice'.",
            "There is a pressing need for an automated, intelligent analysis system that can interpret the semantic meaning of an email and independently evaluate the structural integrity of any embedded URLs. Moreover, end-users require interpretable risk information—not just a binary block/allow decision—to make informed security choices. The absence of an integrated, user-facing system that combines advanced NLP with transparent risk assessment forms the core problem this project addresses."
        ],
        "1.4 Scope": [
            "The scope of this project encompasses the end-to-end development of a full-stack, AI-driven phishing detection platform. In the machine-learning domain, the scope is strictly limited to binary sequence classification using a pre-trained BERT architecture fine-tuned on a composite dataset of verified phishing and legitimate emails.",
            "The URL analysis scope is confined to static, deterministic heuristics (e.g., regex-based extraction, length checking, IP obfuscation detection) and does not extend to dynamic behavioral analysis or live web scraping of the target domain. This ensures high-speed inference without the latency or security risks associated with executing external network requests.",
            "From an application perspective, the scope includes the development of a secure backend API using FastAPI and a responsive, component-based frontend using React. It also covers the integration of a database layer for maintaining analysis history, providing users with a comprehensive dashboard to monitor their threat landscape over time."
        ],
        "1.5 Limitations": [
            "While highly effective at semantic analysis, the current system is constrained by its static approach to URL analysis. Because it does not dynamically resolve domains or check live threat intelligence feeds (such as VirusTotal), it may fail to identify newly registered, clean-looking domains that host malicious content.",
            "Additionally, the BERT model is computationally expensive. Running transformer-based inference requires significant CPU/GPU resources, limiting the system's ability to process tens of thousands of emails per second on standard hardware without significant horizontal scaling and load balancing.",
            "Finally, the model's accuracy is heavily dependent on its training data. Because social engineering tactics evolve rapidly, the model may experience degraded performance against zero-day phishing narratives that deviate significantly from the semantic patterns present in the original training corpus."
        ],
        "1.6 Overview": [
            "The proposed PhishGuard AI system operates through a seamless, decoupled microservices architecture designed for scalability and real-time processing. The workflow begins when a user submits an email—either via manual text input or through the Gmail-integrated inbox—to the React-based frontend. This user interface provides a responsive, intuitive environment for interaction.",
            "Upon submission, the data is transmitted via RESTful API to the FastAPI backend. The backend orchestrates the entire detection pipeline. First, the email text undergoes rigorous preprocessing, including HTML stripping, special character removal, and normalization, preparing it for the machine learning layer.",
            "The core NLP engine utilizes a fine-tuned BERT model. The preprocessed text is tokenized into input IDs and attention masks, then passed through the BERT encoder. The model outputs a continuous probability score representing the likelihood that the text contains phishing semantics. Simultaneously, a URL extraction module scans the text for hyperlinks, passing them to a static analysis engine that evaluates structural heuristics.",
            "Finally, the Risk Assessment Layer synthesizes the BERT probability and the URL indicators using deterministic logic to assign a final risk level (LOW, MEDIUM, or HIGH). This comprehensive evaluation is immediately returned to the frontend for visualization and permanently logged in the SQLite database for historical reporting."
        ],
        "3.1.2 Data Preprocessing": [
            "Data preprocessing is an indispensable stage in NLP, particularly when dealing with raw email corpora which are inherently noisy. The pipeline begins with HTML stripping using specialized parsers like BeautifulSoup, which removes tags, inline CSS, and JavaScript that hold no semantic value for the BERT model.",
            "Subsequently, the system normalizes whitespace, standardizes casing, and removes non-ASCII special characters that could disrupt tokenization. Duplicate removal is strictly enforced by generating MD5 hashes of the canonical text bodies, ensuring that the model does not overfit to repeated spam templates during training.",
            "Finally, the labels are normalized into a strict binary format (0 for Legitimate, 1 for Phishing). The preprocessing rationale centers on the principle that the model should learn to identify intent from human-readable language, not from HTML artifacts or formatting idiosyncrasies."
        ],
        "3.1.4 Canonical Text": [
            "Canonical text preprocessing is a critical prerequisite for transformer-based inference. Once the raw email is stripped of HTML and special characters, the text must be formatted into a tensor representation that the BERT model can ingest. This is achieved using the WordPiece tokenizer associated with the `bert-base-uncased` vocabulary.",
            "The tokenizer splits the normalized text into subword tokens, mapping each to a unique integer ID. To maintain structural consistency across all inputs, the sequence is either padded or truncated to a fixed maximum sequence length of 256 tokens. This length was selected to capture the core intent of the email while optimizing computational efficiency.",
            "In addition to the input IDs, the tokenizer generates an attention mask—a binary tensor indicating which tokens represent actual text and which are merely padding. This mask prevents the BERT encoder from allocating computational attention to empty padding tokens. Finally, special tokens `[CLS]` and `[SEP]` are prepended and appended to the sequence."
        ],
        "3.1.5 BERT": [
            "The machine learning layer is powered by a pre-trained BERT model, specifically fine-tuned for binary sequence classification. Unlike traditional recurrent neural networks, BERT processes the entire sequence simultaneously using the self-attention mechanism, allowing it to evaluate the contextual significance of each word in relation to all other words in the sequence.",
            "The architecture consists of a multi-layer bidirectional Transformer encoder. When the tokenized input IDs and attention masks are fed into the network, they are transformed through token embeddings, segment embeddings, and positional embeddings. These embeddings propagate through the encoder layers, resulting in a rich, high-dimensional contextual representation.",
            "For the classification task, the model extracts the final hidden state corresponding to the `[CLS]` token. This state acts as an aggregate representation of the entire sequence. It is passed through a linear classification head consisting of a fully connected layer and a dropout layer for regularization. The output logits are normalized using a softmax function to produce a probability distribution. A classification threshold of 0.50 determines the preliminary prediction."
        ],
        "3.1.6 URL Extraction": [
            "The URL Extraction and Static Analysis module operates in parallel with the NLP pipeline to evaluate the structural integrity of embedded hyperlinks. A robust regular expression engine parses the raw email body to extract all valid URLs. These URLs are then subjected to a series of static heuristic checks designed to identify common obfuscation techniques.",
            "The static analysis evaluates multiple indicators. It checks for the presence of raw IP addresses in the hostname, which is a strong indicator of malicious intent used to bypass domain reputation filters. It analyzes the protocol, flagging HTTP usage over the secure HTTPS protocol. The length of the URL is assessed, as attackers frequently use excessively long URLs to hide malicious parameters.",
            "Furthermore, the engine counts the number of subdomains, flagging URLs that attempt to mimic legitimate domains via deep subdomain nesting. It also detects the presence of the `@` symbol, which is used to ignore preceding authentication strings, and checks for excessive percent-encoding. The aggregation of these structural anomalies provides critical forensic evidence."
        ],
        "3.1.7 Risk Assessment": [
            "The implementation of the Risk Assessment Layer successfully bridges the gap between raw machine learning probabilities and actionable security intelligence. By defining strict, deterministic rules, the system effectively mitigates the inherent uncertainty of probabilistic models.",
            "The 'LOW' risk category correctly encapsulates emails where the BERT phishing probability falls below the 0.50 threshold and no suspicious URL indicators are triggered. This category effectively represents safe, standard communication.",
            "The 'HIGH' risk category proves highly effective at identifying critical threats. Emails that trigger both a high BERT probability and severe URL heuristic flags (such as raw IP addresses) are immediately escalated to HIGH. Furthermore, the logic successfully catches edge cases where the text is relatively benign, but the URL is blatantly malicious, demonstrating the value of the hybrid evaluation approach."
        ],
        "3.3 Technologies Used": [
            "Python serves as the foundational programming language for the backend, chosen for its unparalleled ecosystem in machine learning and data science. PyTorch provides the underlying deep learning framework, enabling dynamic computational graphs and GPU acceleration for the BERT model, while the Hugging Face Transformers library supplies the pre-trained `bert-base-uncased` architecture and associated WordPiece tokenizer.",
            "The backend REST API is powered by FastAPI, a modern, high-performance web framework for building APIs with Python 3.7+ based on standard Python type hints. Uvicorn acts as the lightning-fast ASGI server, while Pydantic enforces strict data validation on all incoming JSON payloads, preventing malformed requests from crashing the inference pipeline.",
            "On the frontend, React.js dictates the declarative, component-based user interface, allowing for stateful, dynamic rendering of analysis results without page reloads. Vite serves as the ultra-fast build tool, and Tailwind CSS provides the utility-first CSS framework used to rapidly design the modern, responsive, and accessible UI components."
        ],
        "3.4 System Architecture": [
            "The system architecture of PhishGuard AI is designed as a decoupled, multi-tier microservices application, ensuring scalability, maintainability, and clear separation of concerns. The architecture is logically divided into several interconnected layers, each responsible for a specific domain of the detection workflow.",
            "The Presentation Layer, built with React and Tailwind CSS, provides a responsive, single-page application (SPA) interface. It handles all user interactions, input validation, and visual rendering of analysis results. The Authentication Layer secures the system using JWT (JSON Web Tokens), managing user sessions and authorization.",
            "The Backend/API Layer, implemented in FastAPI, serves as the central orchestrator. It exposes RESTful endpoints and manages data flow between the frontend and the analytical engines. The NLP/ML Layer houses the PyTorch-based BERT model, executing the heavy computational tasks required for text inference. Simultaneously, the URL Analysis Layer performs deterministic static checks on extracted hyperlinks.",
            "The Risk Assessment Layer acts as the final decision matrix, synthesizing the probabilistic output from the ML Layer and the heuristic flags from the URL Layer. Finally, the Database/History Layer utilizes SQLite to persistently store user profiles, analysis logs, and system metrics, which are later aggregated and visualized by the Reporting Layer on the frontend dashboard."
        ],
        "4.2 Modules": [
            "The implementation of PhishGuard AI is modularized to ensure high cohesion and low coupling. The Authentication Module handles secure user registration and login, utilizing bcrypt for password hashing and issuing JWTs for stateless session management.",
            "The Gmail Integration Module leverages the Google OAuth2.0 API, allowing users to securely connect their Gmail accounts. It retrieves recent emails and formats them for the Inbox Module, which provides a familiar, mail-client-like interface for users to browse and select emails for scanning.",
            "The core analytical functionality resides in the Email Analysis Module. It accepts raw email text, sanitizes it through the Data Preprocessing Module, and delegates the computational workload to the BERT Inference Module and URL Analysis Module. The Risk Assessment Module then synthesizes these discrete outputs into a unified risk profile.",
            "Post-analysis, the Reports and History Modules capture the transaction. The History Module writes the metadata, predictions, and extracted indicators to the relational database. The Dashboard Module queries this historical data to render real-time statistical visualizations, providing users with actionable intelligence regarding their threat landscape."
        ],
        "7.5 Future Enhancements": [
            "While PhishGuard AI successfully demonstrates the viability of transformer-based phishing detection, several avenues exist for future enhancement. Improved model training utilizing larger, more diverse, and continually updated datasets would further increase precision and adaptability against zero-day social engineering tactics.",
            "The URL intelligence module could be significantly upgraded by integrating dynamic analysis and real-time threat intelligence feeds (e.g., VirusTotal or Google Safe Browsing API), moving beyond static heuristics to verify domain reputation and active malicious payloads.",
            "Additionally, advancements in model explainability, such as integrating SHAP or LIME, would provide users with precise highlighting of the specific words or phrases that triggered the phishing classification. Future deployment strategies involve migrating the architecture to a containerized cloud environment using Docker and Kubernetes, ensuring high availability and enterprise-grade scalability."
        ]
    }
    
    body = doc._body._body
    elements_to_remove = []
    
    for element in body.iterchildren():
        if element.tag.endswith('p'):
            from docx.text.paragraph import Paragraph
            p = Paragraph(element, doc)
            text = p.text.strip()
            
            # Remove all boilerplate paragraphs
            if text.startswith("The purpose of ") and "is integral to establishing the core technical foundations" in text:
                elements_to_remove.append(element)
    
    # Actually remove them
    for el in elements_to_remove:
        try:
            el.getparent().remove(el)
        except Exception:
            pass

    # Now inject the expansions using 'in' instead of exact match
    paragraphs = doc.paragraphs
    insert_count = 0
    
    for i in range(len(paragraphs) - 1, -1, -1):
        p = paragraphs[i]
        text = p.text.strip()
        
        for key, paras_to_insert in expansion_dict.items():
            if key in text and len(text) < 100: # Make sure it's a heading, not a body paragraph
                for new_text in reversed(paras_to_insert):
                    new_p = doc.add_paragraph(new_text)
                    new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    for run in new_p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(12)
                    
                    p._element.addnext(new_p._element)
                    insert_count += 1
                
                del expansion_dict[key]
                break

    print(f"Removed {len(elements_to_remove)} boilerplate paragraphs.")
    print(f"Successfully injected {insert_count} highly detailed paragraphs.")
    
    temp_path = r"D:\AICS PROJECT\Project_Report_PhishGuard_AI_temp.docx"
    doc.save(temp_path)
    print("Report thoroughly cleaned, expanded, and saved successfully.")

if __name__ == "__main__":
    main()
