import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = None
        if level == 1:
            run.font.size = Pt(14)
            run.font.bold = True
        else:
            run.font.size = Pt(12)
            run.font.bold = True

def add_paragraph(doc, text, bold=False, italic=False, align="justify"):
    p = doc.add_paragraph()
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.25
    return p

def create_table(doc, headers, rows, title=""):
    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        run.bold = True
        p.paragraph_format.space_after = Pt(6)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.bold = True
                r.font.size = Pt(11)
    
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = str(cell_data)
            for p in row_cells[i].paragraphs:
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(11)
    
    doc.add_paragraph("") # Space after table

def generate_report():
    doc = Document()
    
    # 1. Title Page
    for _ in range(5): doc.add_paragraph()
    add_paragraph(doc, "PROJECT REPORT ON", bold=True, align="center")
    add_paragraph(doc, "AI-BASED EMAIL PHISHING DETECTION USING BERT", bold=True, align="center")
    doc.add_page_break()

    add_heading(doc, "CHAPTER 3 – METHODOLOGY", level=1)
    add_heading(doc, "3.1 System Overview", level=2)
    add_heading(doc, "3.1.1 Dataset Collection", level=3)
    add_paragraph(doc, "The dataset was collected from verified sources like CEAS, Enron, and Ling-Spam.")
    create_table(doc, ["Dataset Source", "Description", "Role in Project"], [
        ["CEAS", "Conference on Email and Anti-Spam", "Spam/phishing examples"],
        ["Enron", "Enron email corpus", "Baseline legitimate"],
        ["Ling-Spam", "Linguistics mailing list", "Academic spam structures"]
    ], title="Table 3.1 — Dataset Sources and Description")

    add_heading(doc, "3.1.2 Data Preprocessing and Cleaning", level=3)
    add_paragraph(doc, "Data cleaning involved removing HTML and duplicates.")
    create_table(doc, ["Processing Stage", "Description", "Result"], [
        ["Data Cleaning", "Remove HTML and special chars", "Cleaned text sequences"],
        ["Duplicate Removal", "Identify exact duplicates", "13,250 duplicates removed"]
    ], title="Table 3.3 — Data Preprocessing Summary")
    
    add_heading(doc, "3.1.3 Data Splitting", level=3)
    create_table(doc, ["Dataset", "Number of Samples", "Percentage"], [
        ["Training", "70,308", "80%"],
        ["Validation", "8,789", "10%"],
        ["Testing", "8,789", "10%"],
        ["Total", "87,886", "100%"]
    ], title="Table 3.4 — Train, Validation and Test Dataset Distribution")

    add_heading(doc, "3.1.5 BERT-Based Email Classification", level=3)
    create_table(doc, ["Parameter", "Configuration"], [
        ["Model", "bert-base-uncased"],
        ["Architecture", "BERT Sequence Classification"],
        ["Number of Classes", "2"],
        ["Class 0", "Legitimate"],
        ["Class 1", "Phishing"],
        ["Maximum Sequence Length", "256"],
        ["Tokenizer", "BERT Tokenizer"],
        ["Framework", "PyTorch"],
        ["Classification Threshold", "0.50"]
    ], title="Table 3.5 — BERT Model Configuration")
    
    add_heading(doc, "3.3 Technologies Used", level=2)
    create_table(doc, ["Layer", "Technology", "Purpose"], [
        ["Frontend", "React, Vite, Tailwind CSS", "User Interface"],
        ["Backend", "FastAPI, Uvicorn, Python", "RESTful API Server"],
        ["Machine Learning", "PyTorch, Hugging Face, BERT", "Text Inference"],
        ["Data Processing", "Pandas", "Dataset manipulation"]
    ], title="Table 3.6 — Technology Stack")
    
    add_heading(doc, "3.1.7 Risk Assessment", level=3)
    create_table(doc, ["Risk Level", "BERT/URL Conditions", "Interpretation"], [
        ["LOW", "Low phishing probability and limited suspicious indicators", "Lower observed risk"],
        ["MEDIUM", "Intermediate probability or URL indicators", "Requires attention"],
        ["HIGH", "High phishing probability and/or strong URL indicators", "Higher observed risk"]
    ], title="Table 3.7 — Risk Assessment Levels")

    doc.add_page_break()
    add_heading(doc, "CHAPTER 4 – IMPLEMENTATION", level=1)
    create_table(doc, ["Module", "Main Function", "Implementation"], [
        ["Authentication", "Secure access", "JWT Tokens"],
        ["Email Analysis", "Text scanning", "BERT Pipeline"],
        ["Risk Assessment", "Threshold calculation", "Logical gates"]
    ], title="Table 4.1 — System Modules and Functions")
    
    create_table(doc, ["Method", "Endpoint", "Purpose"], [
        ["GET", "/health", "API Status"],
        ["GET", "/model-info", "Model Metadata"],
        ["POST", "/predict", "Main Inference Endpoint"]
    ], title="Table 4.2 — API Endpoints")
    
    create_table(doc, ["Page", "Route", "Main Function"], [
        ["Login", "/login", "Authentication"],
        ["Dashboard", "/dashboard", "Metrics view"],
        ["Analyze", "/analyze", "Manual email scan"]
    ], title="Table 4.3 — Application Pages and Functions")

    doc.add_page_break()
    add_heading(doc, "CHAPTER 5 – TESTING", level=1)
    create_table(doc, ["Component", "Configuration"], [
        ["Operating System", "Windows 11 / Ubuntu 22.04"],
        ["Python", "3.10"],
        ["Frameworks", "FastAPI, React"]
    ], title="Table 5.1 — Test Environment and Configuration")
    
    create_table(doc, ["Test ID", "Module", "Test Description", "Expected Result", "Actual Result", "Status"], [
        ["TC01", "API", "Test /health endpoint", "200 OK", "200 OK", "PASS"],
        ["TC02", "ML", "Verify valid payload inference", "Probabilities returned", "Probabilities returned", "PASS"]
    ], title="Table 5.2 — Test Cases and Results")

    doc.add_page_break()
    add_heading(doc, "CHAPTER 6 – RESULTS", level=1)
    add_heading(doc, "6.1.3 Prediction Performance Analysis", level=3)
    create_table(doc, ["Metric", "Value"], [
        ["Accuracy", "67.58%"],
        ["Precision", "61.28%"],
        ["Recall", "97.48%"],
        ["F1-score", "75.25%"]
    ], title="Table 6.1 — Model Evaluation Metrics (Documented Historical Results)")
    
    create_table(doc, ["Email Type", "BERT Prediction", "Phishing Probability", "URL Indicators", "Risk Level"], [
        ["Password Reset", "Phishing", "99.1%", "IP Domain", "HIGH"],
        ["Newsletter", "Legitimate", "12.4%", "None", "LOW"]
    ], title="Table 6.2 — Sample Email Analysis Results")

    doc.add_page_break()
    add_heading(doc, "REFERENCES", level=1)
    add_paragraph(doc, "[1] J. Devlin, M. Chang, K. Lee, and K. Toutanova, 'BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding', arXiv preprint arXiv:1810.04805, 2018.")
    add_paragraph(doc, "[2] A. Vaswani et al., 'Attention Is All You Need', Advances in Neural Information Processing Systems, 2017.")
    add_paragraph(doc, "[3] FastAPI Documentation, 'FastAPI framework, high performance, easy to learn, fast to code, ready for production'. https://fastapi.tiangolo.com/")

    doc.save("Project_Report_PhishGuard_AI_Final.docx")
    print("Report v4 generated successfully with strict table adherence.")

if __name__ == "__main__":
    generate_report()
